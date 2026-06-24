"""Excel 报表生成工具类"""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import List, Dict, Any
from core.config import settings
from datetime import datetime


def generate_report(data_rows: List[Dict[str, Any]], columns: List[str], report_name: str) -> str:
    """
    生成标准化业务报表 Excel
    :param data_rows: 清洗后的数据行列表，每行为 dict
    :param columns: 列名列表
    :param report_name: 报表名称
    :return: 生成的 Excel 文件路径
    """
    wb = Workbook()
    ws = wb.active
    ws.title = report_name[:31]

    # 样式
    header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1677FF", end_color="1677FF", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    cell_font = Font(name="微软雅黑", size=10)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # 写表头
    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # 写数据
    for row_idx, row_data in enumerate(data_rows, 2):
        for col_idx, col_name in enumerate(columns, 1):
            value = row_data.get(col_name, "")
            if isinstance(value, datetime):
                value = value.strftime("%Y-%m-%d %H:%M:%S")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = cell_font
            cell.alignment = Alignment(vertical="center")
            cell.border = thin_border

    # 自适应列宽
    for col_idx, col_name in enumerate(columns, 1):
        max_len = len(col_name) * 2
        for row_idx in range(2, len(data_rows) + 2):
            cell_val = ws.cell(row=row_idx, column=col_idx).value
            if cell_val:
                max_len = max(max_len, len(str(cell_val)) * 1.5)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 60)

    # 保存
    from utils.common import ensure_dir
    report_dir = os.path.join(settings.EXCEL_STORAGE_DIR, "reports")
    ensure_dir(report_dir)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    file_path = os.path.join(report_dir, f"{report_name}_{timestamp}.xlsx")
    wb.save(file_path)
    return file_path
